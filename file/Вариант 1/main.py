"""GUI приложение для поиска организаций без сайтов на Яндекс.Картах."""
import os
import sys

# При запуске из .exe (PyInstaller) указываем путь к встроенному Chromium
if getattr(sys, "frozen", False):
    _bundled = os.path.join(getattr(sys, "_MEIPASS", ""), "ms-playwright")
    if os.path.isdir(_bundled):
        os.environ["PLAYWRIGHT_BROWSERS_PATH"] = _bundled

import asyncio
import threading
from datetime import datetime
from tkinter import ttk, StringVar, IntVar, BooleanVar, END, filedialog, messagebox

import customtkinter as ctk

from parser import search_async


# ===== Цветовая палитра =====
BG_DARK = "#0f1419"
BG_PANEL = "#1a1f29"
BG_CARD = "#222834"
BG_INPUT = "#2a3140"
BG_HOVER = "#323a4a"
ACCENT = "#3b82f6"
ACCENT_HOVER = "#2563eb"
ACCENT_GREEN = "#10b981"
ACCENT_GREEN_HOVER = "#059669"
ACCENT_RED = "#ef4444"
ACCENT_RED_HOVER = "#dc2626"
TEXT_PRIMARY = "#e6edf3"
TEXT_SECONDARY = "#8b949e"
TEXT_MUTED = "#6e7681"
BORDER = "#30363d"

SPHERES = [
    "Кафе", "Ресторан", "Кофейня", "Пиццерия", "Бар", "Столовая",
    "Барбершоп", "Парикмахерская", "Салон красоты", "Маникюрный салон",
    "Стоматология", "Медицинский центр", "Аптека", "Ветеринарная клиника",
    "Автосервис", "Автомойка", "Шиномонтаж",
    "Магазин одежды", "Цветочный магазин", "Зоомагазин",
    "Фитнес-клуб", "Йога-студия", "Танцевальная студия",
    "Юридические услуги", "Бухгалтерские услуги", "Агентство недвижимости",
    "Ремонт квартир", "Дизайн интерьера", "Архитектурное бюро",
    "Хостел", "Гостиница", "Турагентство",
    "Языковая школа", "Детский центр", "Автошкола",
]

CITIES = [
    "Москва", "Санкт-Петербург", "Новосибирск", "Екатеринбург", "Казань",
    "Нижний Новгород", "Челябинск", "Самара", "Уфа", "Ростов-на-Дону",
    "Краснодар", "Омск", "Воронеж", "Пермь", "Волгоград",
    "Сочи", "Геленджик", "Анапа", "Новороссийск", "Калининград",
    "Тюмень", "Ижевск", "Барнаул", "Ульяновск", "Иркутск",
    "Хабаровск", "Владивосток", "Ярославль", "Махачкала", "Томск",
    "Оренбург", "Кемерово", "Рязань", "Тольятти", "Астрахань",
    "Пенза", "Липецк", "Тула", "Киров", "Чебоксары",
    "Калуга", "Брянск", "Курск", "Тверь", "Ставрополь",
    "Иваново", "Магнитогорск", "Сургут", "Белгород", "Владимир",
]


class ParserApp:
    def __init__(self):
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.root = ctk.CTk()
        self.root.title("Яндекс.Карты — Поиск организаций без сайта")
        self.root.geometry("1280x820")
        self.root.minsize(1080, 680)
        self.root.configure(fg_color=BG_DARK)

        self.worker_thread = None
        self.stop_requested = False
        self.results = []

        self._setup_treeview_style()
        self._build_ui()

    # ---------- Стиль таблицы ----------
    def _setup_treeview_style(self):
        style = ttk.Style()
        style.theme_use("default")
        style.configure(
            "Modern.Treeview",
            background=BG_PANEL,
            foreground=TEXT_PRIMARY,
            fieldbackground=BG_PANEL,
            rowheight=34,
            borderwidth=0,
            font=("Segoe UI", 10),
        )
        style.configure(
            "Modern.Treeview.Heading",
            background=BG_CARD,
            foreground=TEXT_SECONDARY,
            font=("Segoe UI Semibold", 10),
            borderwidth=0,
            relief="flat",
            padding=(10, 10),
        )
        style.map(
            "Modern.Treeview",
            background=[("selected", ACCENT)],
            foreground=[("selected", "white")],
        )
        style.map(
            "Modern.Treeview.Heading",
            background=[("active", BG_HOVER)],
        )
        style.layout(
            "Modern.Treeview",
            [("Modern.Treeview.treearea", {"sticky": "nswe"})],
        )

    # ---------- Сборка UI ----------
    def _build_ui(self):
        # Шапка
        self._build_header()

        # Основной контейнер с двумя колонками
        body = ctk.CTkFrame(self.root, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=20, pady=(0, 20))
        body.grid_columnconfigure(0, weight=0, minsize=320)
        body.grid_columnconfigure(1, weight=1)
        body.grid_rowconfigure(0, weight=1)

        self._build_sidebar(body)
        self._build_main(body)

    def _build_header(self):
        header = ctk.CTkFrame(self.root, fg_color="transparent", height=80)
        header.pack(fill="x", padx=20, pady=(20, 10))
        header.pack_propagate(False)

        icon_box = ctk.CTkFrame(
            header, width=56, height=56, corner_radius=14,
            fg_color=ACCENT,
        )
        icon_box.pack(side="left", padx=(0, 16))
        icon_box.pack_propagate(False)
        ctk.CTkLabel(
            icon_box, text="📍", font=ctk.CTkFont(size=28),
            text_color="white",
        ).pack(expand=True)

        text_frame = ctk.CTkFrame(header, fg_color="transparent")
        text_frame.pack(side="left", fill="y")
        ctk.CTkLabel(
            text_frame, text="Поиск организаций",
            font=ctk.CTkFont(family="Segoe UI", size=22, weight="bold"),
            text_color=TEXT_PRIMARY,
        ).pack(anchor="w")
        ctk.CTkLabel(
            text_frame, text="Лидогенерация: компании без сайтов на Яндекс.Картах",
            font=ctk.CTkFont(family="Segoe UI", size=12),
            text_color=TEXT_SECONDARY,
        ).pack(anchor="w")

        # Индикатор статуса
        self.status_dot = ctk.CTkLabel(
            header, text="●", font=ctk.CTkFont(size=18),
            text_color=TEXT_MUTED,
        )
        self.status_dot.pack(side="right", padx=(0, 6))
        self.status_text = ctk.CTkLabel(
            header, text="Готов к работе",
            font=ctk.CTkFont(family="Segoe UI", size=12),
            text_color=TEXT_SECONDARY,
        )
        self.status_text.pack(side="right", padx=(0, 8))

    # ---------- Сайдбар ----------
    def _build_sidebar(self, parent):
        sidebar = ctk.CTkFrame(parent, fg_color=BG_PANEL, corner_radius=16)
        sidebar.grid(row=0, column=0, sticky="nsew", padx=(0, 16))

        inner = ctk.CTkFrame(sidebar, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=20, pady=20)

        # --- Параметры поиска ---
        self._section_label(inner, "ПАРАМЕТРЫ ПОИСКА")

        self._field_label(inner, "Сфера деятельности")
        self.sphere_var = StringVar(value=SPHERES[0])
        sphere_cb = ctk.CTkComboBox(
            inner, values=SPHERES, variable=self.sphere_var,
            height=38, corner_radius=10,
            fg_color=BG_INPUT, border_color=BORDER, border_width=1,
            button_color=BG_INPUT, button_hover_color=BG_HOVER,
            dropdown_fg_color=BG_CARD, dropdown_hover_color=BG_HOVER,
            text_color=TEXT_PRIMARY, dropdown_text_color=TEXT_PRIMARY,
            font=ctk.CTkFont(family="Segoe UI", size=12),
        )
        sphere_cb.pack(fill="x", pady=(0, 14))

        self._field_label(inner, "Город")
        self.city_var = StringVar(value="Геленджик")
        city_cb = ctk.CTkComboBox(
            inner, values=CITIES, variable=self.city_var,
            height=38, corner_radius=10,
            fg_color=BG_INPUT, border_color=BORDER, border_width=1,
            button_color=BG_INPUT, button_hover_color=BG_HOVER,
            dropdown_fg_color=BG_CARD, dropdown_hover_color=BG_HOVER,
            text_color=TEXT_PRIMARY, dropdown_text_color=TEXT_PRIMARY,
            font=ctk.CTkFont(family="Segoe UI", size=12),
        )
        city_cb.pack(fill="x", pady=(0, 14))

        self._field_label(inner, "Максимум результатов")
        self.max_var = IntVar(value=60)
        max_entry = ctk.CTkEntry(
            inner, textvariable=self.max_var,
            height=38, corner_radius=10,
            fg_color=BG_INPUT, border_color=BORDER, border_width=1,
            text_color=TEXT_PRIMARY,
            font=ctk.CTkFont(family="Segoe UI", size=12),
        )
        max_entry.pack(fill="x", pady=(0, 18))

        # --- Опции ---
        self._section_label(inner, "ОПЦИИ")
        self.only_no_site = BooleanVar(value=True)
        ctk.CTkCheckBox(
            inner, text="Только без сайта", variable=self.only_no_site,
            fg_color=ACCENT, hover_color=ACCENT_HOVER,
            text_color=TEXT_PRIMARY, border_color=BORDER,
            font=ctk.CTkFont(family="Segoe UI", size=12),
            corner_radius=5,
        ).pack(anchor="w", pady=(0, 8))

        self.headless_var = BooleanVar(value=True)
        ctk.CTkCheckBox(
            inner, text="Скрытый браузер", variable=self.headless_var,
            fg_color=ACCENT, hover_color=ACCENT_HOVER,
            text_color=TEXT_PRIMARY, border_color=BORDER,
            font=ctk.CTkFont(family="Segoe UI", size=12),
            corner_radius=5,
        ).pack(anchor="w", pady=(0, 20))

        # --- Кнопки ---
        self.start_btn = ctk.CTkButton(
            inner, text="🔍   Запустить поиск", command=self.start_search,
            height=44, corner_radius=10,
            fg_color=ACCENT, hover_color=ACCENT_HOVER,
            text_color="white",
            font=ctk.CTkFont(family="Segoe UI", size=13, weight="bold"),
        )
        self.start_btn.pack(fill="x", pady=(0, 8))

        self.stop_btn = ctk.CTkButton(
            inner, text="⏹   Остановить", command=self.stop_search,
            height=38, corner_radius=10,
            fg_color="transparent", hover_color=BG_CARD,
            border_color=ACCENT_RED, border_width=1,
            text_color=ACCENT_RED,
            font=ctk.CTkFont(family="Segoe UI", size=12),
            state="disabled",
        )
        self.stop_btn.pack(fill="x", pady=(0, 16))

        # Разделитель
        ctk.CTkFrame(inner, height=1, fg_color=BORDER).pack(fill="x", pady=10)

        self.export_btn = ctk.CTkButton(
            inner, text="📥   Экспорт в Excel", command=self.export_excel,
            height=40, corner_radius=10,
            fg_color=ACCENT_GREEN, hover_color=ACCENT_GREEN_HOVER,
            text_color="white",
            font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
            state="disabled",
        )
        self.export_btn.pack(fill="x", side="bottom")

    def _section_label(self, parent, text):
        ctk.CTkLabel(
            parent, text=text,
            font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold"),
            text_color=TEXT_MUTED,
        ).pack(anchor="w", pady=(0, 10))

    def _field_label(self, parent, text):
        ctk.CTkLabel(
            parent, text=text,
            font=ctk.CTkFont(family="Segoe UI", size=11),
            text_color=TEXT_SECONDARY,
        ).pack(anchor="w", pady=(0, 4))

    # ---------- Основная область ----------
    def _build_main(self, parent):
        main = ctk.CTkFrame(parent, fg_color="transparent")
        main.grid(row=0, column=1, sticky="nsew")
        main.grid_rowconfigure(1, weight=1)
        main.grid_rowconfigure(3, weight=0)
        main.grid_columnconfigure(0, weight=1)

        # Карточки статистики
        stats = ctk.CTkFrame(main, fg_color="transparent")
        stats.grid(row=0, column=0, sticky="ew", pady=(0, 14))
        stats.grid_columnconfigure((0, 1, 2), weight=1)

        self.stat_found = self._stat_card(stats, "Найдено", "0", ACCENT_GREEN, 0)
        self.stat_processed = self._stat_card(stats, "Обработано", "0", ACCENT, 1)
        self.stat_query = self._stat_card(stats, "Текущий запрос", "—", "#a78bfa", 2)

        # Таблица
        table_card = ctk.CTkFrame(main, fg_color=BG_PANEL, corner_radius=16)
        table_card.grid(row=1, column=0, sticky="nsew")

        table_inner = ctk.CTkFrame(table_card, fg_color="transparent")
        table_inner.pack(fill="both", expand=True, padx=16, pady=16)

        ctk.CTkLabel(
            table_inner, text="Результаты",
            font=ctk.CTkFont(family="Segoe UI", size=14, weight="bold"),
            text_color=TEXT_PRIMARY,
        ).pack(anchor="w", pady=(0, 10))

        tree_wrap = ctk.CTkFrame(table_inner, fg_color=BG_PANEL, corner_radius=8)
        tree_wrap.pack(fill="both", expand=True)

        cols = ("idx", "name", "address", "phone", "site", "category", "city")
        self.tree = ttk.Treeview(
            tree_wrap, columns=cols, show="headings",
            style="Modern.Treeview", height=12,
        )
        headers = {
            "idx": ("№", 50),
            "name": ("Название", 240),
            "address": ("Адрес", 260),
            "phone": ("Телефон", 130),
            "site": ("Сайт", 150),
            "category": ("Сфера", 120),
            "city": ("Город", 110),
        }
        for col, (title, w) in headers.items():
            self.tree.heading(col, text=title)
            anchor = "center" if col == "idx" else "w"
            self.tree.column(col, width=w, anchor=anchor)

        self.tree.tag_configure("odd", background=BG_PANEL)
        self.tree.tag_configure("even", background="#1d2330")

        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        # Прогресс
        prog_frame = ctk.CTkFrame(main, fg_color="transparent")
        prog_frame.grid(row=2, column=0, sticky="ew", pady=(14, 10))
        self.progress = ctk.CTkProgressBar(
            prog_frame, height=6, corner_radius=3,
            progress_color=ACCENT, fg_color=BG_CARD,
        )
        self.progress.pack(fill="x")
        self.progress.set(0)

        # Лог
        log_card = ctk.CTkFrame(main, fg_color=BG_PANEL, corner_radius=16, height=180)
        log_card.grid(row=3, column=0, sticky="ew")
        log_card.grid_propagate(False)

        log_inner = ctk.CTkFrame(log_card, fg_color="transparent")
        log_inner.pack(fill="both", expand=True, padx=16, pady=14)

        log_header = ctk.CTkFrame(log_inner, fg_color="transparent")
        log_header.pack(fill="x", pady=(0, 8))
        ctk.CTkLabel(
            log_header, text="Журнал",
            font=ctk.CTkFont(family="Segoe UI", size=13, weight="bold"),
            text_color=TEXT_PRIMARY,
        ).pack(side="left")
        ctk.CTkButton(
            log_header, text="Очистить", command=self._clear_log,
            width=80, height=24, corner_radius=6,
            fg_color="transparent", hover_color=BG_CARD,
            border_color=BORDER, border_width=1,
            text_color=TEXT_SECONDARY,
            font=ctk.CTkFont(family="Segoe UI", size=10),
        ).pack(side="right")

        self.log = ctk.CTkTextbox(
            log_inner,
            fg_color=BG_DARK, text_color=TEXT_SECONDARY,
            font=ctk.CTkFont(family="Consolas", size=10),
            corner_radius=8, border_width=0,
            wrap="word",
        )
        self.log.pack(fill="both", expand=True)

    def _stat_card(self, parent, title, value, accent, col):
        card = ctk.CTkFrame(parent, fg_color=BG_PANEL, corner_radius=14, height=84)
        card.grid(row=0, column=col, sticky="ew", padx=(0 if col == 0 else 6, 6 if col < 2 else 0))
        card.grid_propagate(False)

        bar = ctk.CTkFrame(card, fg_color=accent, width=4, corner_radius=2)
        bar.place(x=10, y=14, relheight=0.7)

        ctk.CTkLabel(
            card, text=title,
            font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold"),
            text_color=TEXT_MUTED,
        ).place(x=24, y=14)
        val_label = ctk.CTkLabel(
            card, text=value,
            font=ctk.CTkFont(family="Segoe UI", size=22, weight="bold"),
            text_color=TEXT_PRIMARY,
        )
        val_label.place(x=22, y=34)
        return val_label

    # ---------- Поток ----------
    def start_search(self):
        if self.worker_thread and self.worker_thread.is_alive():
            return
        sphere = self.sphere_var.get().strip()
        city = self.city_var.get().strip()
        if not sphere or not city:
            messagebox.showwarning("Пустые поля", "Укажи сферу и город.")
            return
        try:
            max_n = int(self.max_var.get())
        except Exception:
            max_n = 60

        self.results.clear()
        self.tree.delete(*self.tree.get_children())
        self._clear_log()
        self.stop_requested = False
        self.start_btn.configure(state="disabled", text="🔍   Поиск идёт...")
        self.stop_btn.configure(state="normal")
        self.export_btn.configure(state="disabled")
        self.progress.configure(mode="indeterminate")
        self.progress.start()
        self.status_dot.configure(text_color=ACCENT)
        self.status_text.configure(text="Идёт поиск...")
        self._update_stat(self.stat_found, "0")
        self._update_stat(self.stat_processed, "0")
        self._update_stat(self.stat_query, f"{sphere} · {city}")

        params = dict(
            category=sphere,
            city=city,
            max_results=max_n,
            headless=self.headless_var.get(),
            only_without_site=self.only_no_site.get(),
        )
        self.worker_thread = threading.Thread(target=self._run_worker, args=(params,), daemon=True)
        self.worker_thread.start()

    def stop_search(self):
        self.stop_requested = True
        self._log_safe("⏹ Получен сигнал остановки...")
        self.stop_btn.configure(state="disabled")

    def _run_worker(self, params):
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            loop.run_until_complete(
                search_async(
                    log=self._log_safe,
                    on_result=self._add_row_safe,
                    stop_flag=lambda: self.stop_requested,
                    **params,
                )
            )
        except Exception as e:
            self._log_safe(f"❌ ОШИБКА: {e}")
        finally:
            loop.close()
            self.root.after(0, self._finish)

    def _finish(self):
        self.progress.stop()
        self.progress.configure(mode="determinate")
        self.progress.set(1.0 if self.results else 0)
        self.start_btn.configure(state="normal", text="🔍   Запустить поиск")
        self.stop_btn.configure(state="disabled")
        if self.results:
            self.export_btn.configure(state="normal")
            self.status_dot.configure(text_color=ACCENT_GREEN)
            self.status_text.configure(text=f"Готово · найдено {len(self.results)}")
        else:
            self.status_dot.configure(text_color=TEXT_MUTED)
            self.status_text.configure(text="Готов к работе")

    # ---------- Поток-сейф ----------
    def _log_safe(self, msg):
        self.root.after(0, lambda: self._append_log(msg))

    def _append_log(self, msg):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log.insert(END, f"[{timestamp}]  {msg}\n")
        self.log.see(END)
        # Обновляем счётчик обработанных
        if msg.startswith("[") and "/" in msg:
            try:
                cur = msg.split("[")[1].split("/")[0]
                self._update_stat(self.stat_processed, cur)
            except Exception:
                pass

    def _clear_log(self):
        self.log.delete("1.0", END)

    def _add_row_safe(self, data):
        self.root.after(0, lambda: self._add_row(data))

    def _add_row(self, data):
        idx = len(self.results) + 1
        self.results.append(data)
        tag = "even" if idx % 2 == 0 else "odd"
        self.tree.insert(
            "", END,
            values=(
                idx,
                data.get("name", ""),
                data.get("address", ""),
                data.get("phone", "") or "—",
                data.get("site", "") or "—",
                data.get("category", ""),
                data.get("city", ""),
            ),
            tags=(tag,),
        )
        self._update_stat(self.stat_found, str(idx))

    def _update_stat(self, label, value):
        label.configure(text=value)

    # ---------- Экспорт ----------
    def export_excel(self):
        if not self.results:
            return
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        default_name = (
            f"yandex_{self.sphere_var.get()}_{self.city_var.get()}_"
            f"{datetime.now():%Y%m%d_%H%M}.xlsx"
        ).replace(" ", "_")
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=default_name,
        )
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Организации"

        headers = ["№", "Название", "Адрес", "Телефон", "Сайт", "Сфера", "Город"]
        ws.append(headers)

        header_fill = PatternFill("solid", fgColor="3B82F6")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin = Side(border_style="thin", color="E5E7EB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        ws.row_dimensions[1].height = 28

        for i, r in enumerate(self.results, 1):
            ws.append([
                i,
                r.get("name", ""),
                r.get("address", ""),
                r.get("phone", ""),
                r.get("site", ""),
                r.get("category", ""),
                r.get("city", ""),
            ])
            if i % 2 == 0:
                for cell in ws[i + 1]:
                    cell.fill = PatternFill("solid", fgColor="F3F4F6")

        widths = [6, 38, 42, 20, 28, 22, 18]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        ws.freeze_panes = "A2"

        try:
            wb.save(path)
            messagebox.showinfo("Готово", f"Сохранено:\n{path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить: {e}")

    def run(self):
        self.root.mainloop()


def main():
    app = ParserApp()
    app.run()


if __name__ == "__main__":
    main()
